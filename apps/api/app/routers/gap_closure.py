from __future__ import annotations

import importlib
import re
import sqlite3
import uuid
from datetime import datetime, timedelta
from pathlib import Path
from typing import Any, Optional

from fastapi import APIRouter, File, Form, HTTPException, Query, UploadFile
from fastapi.responses import FileResponse
from pydantic import BaseModel, Field


router = APIRouter(prefix="/gap-closure", tags=["Gap Closure"])


def get_database_path() -> Path:
    try:
        db_module = importlib.import_module("app.services.db")
        for attr in ["DB_PATH", "DATABASE_PATH", "DATABASE_FILE", "SQLITE_PATH"]:
            value = getattr(db_module, attr, None)
            if value:
                return Path(value)
    except Exception:
        pass

    return Path(__file__).resolve().parents[2] / "munjiz.db"


def get_project_root() -> Path:
    return Path(__file__).resolve().parents[2]


def get_evidence_storage_dir() -> Path:
    evidence_dir = get_project_root() / "storage" / "gap_evidence"
    evidence_dir.mkdir(parents=True, exist_ok=True)
    return evidence_dir


def get_connection() -> sqlite3.Connection:
    db_path = get_database_path()
    db_path.parent.mkdir(parents=True, exist_ok=True)
    conn = sqlite3.connect(str(db_path))
    conn.row_factory = sqlite3.Row
    return conn


def now_iso() -> str:
    return datetime.utcnow().replace(microsecond=0).isoformat() + "Z"


def row_to_dict(row: sqlite3.Row | None) -> dict[str, Any] | None:
    return dict(row) if row is not None else None


def rows_to_list(rows: list[sqlite3.Row]) -> list[dict[str, Any]]:
    return [dict(row) for row in rows]


def table_exists(conn: sqlite3.Connection, table_name: str) -> bool:
    row = conn.execute(
        "SELECT name FROM sqlite_master WHERE type='table' AND name = ?",
        (table_name,),
    ).fetchone()
    return row is not None


def column_exists(conn: sqlite3.Connection, table_name: str, column_name: str) -> bool:
    if not table_exists(conn, table_name):
        return False
    rows = conn.execute(f"PRAGMA table_info({table_name})").fetchall()
    return any(row["name"] == column_name for row in rows)


def ensure_column(conn: sqlite3.Connection, table_name: str, column_name: str, column_sql: str) -> None:
    if not column_exists(conn, table_name, column_name):
        conn.execute(f"ALTER TABLE {table_name} ADD COLUMN {column_sql}")


def ensure_gap_tables() -> None:
    with get_connection() as conn:
        conn.execute(
            """
            CREATE TABLE IF NOT EXISTS team_members (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                name TEXT NOT NULL,
                role TEXT NOT NULL UNIQUE,
                title TEXT NOT NULL,
                department TEXT NOT NULL,
                email TEXT,
                active INTEGER NOT NULL DEFAULT 1,
                created_at TEXT NOT NULL
            )
            """
        )

        conn.execute(
            """
            CREATE TABLE IF NOT EXISTS gap_tasks (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                tender_id INTEGER NOT NULL,
                requirement_id INTEGER,
                requirement_title TEXT NOT NULL,
                title TEXT NOT NULL,
                description TEXT NOT NULL,
                owner_role TEXT NOT NULL,
                owner_name TEXT NOT NULL,
                owner_department TEXT NOT NULL,
                priority TEXT NOT NULL,
                status TEXT NOT NULL,
                evidence_type TEXT NOT NULL,
                evidence_instruction TEXT NOT NULL,
                impact_score INTEGER NOT NULL DEFAULT 5,
                due_date TEXT,
                verification_status TEXT NOT NULL DEFAULT 'NOT_SUBMITTED',
                verification_notes TEXT,
                evidence_note TEXT,
                created_by TEXT NOT NULL DEFAULT 'Gap Assignment Agent',
                created_at TEXT NOT NULL,
                updated_at TEXT NOT NULL,
                closed_at TEXT
            )
            """
        )

        conn.execute(
            """
            CREATE TABLE IF NOT EXISTS gap_task_events (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                task_id INTEGER NOT NULL,
                event_type TEXT NOT NULL,
                message TEXT NOT NULL,
                created_at TEXT NOT NULL
            )
            """
        )

        conn.execute(
            """
            CREATE TABLE IF NOT EXISTS task_evidence (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                task_id INTEGER NOT NULL,
                tender_id INTEGER NOT NULL,
                original_filename TEXT NOT NULL,
                stored_filename TEXT NOT NULL,
                file_path TEXT NOT NULL,
                mime_type TEXT,
                file_size INTEGER NOT NULL DEFAULT 0,
                evidence_note TEXT,
                extracted_text TEXT,
                extraction_status TEXT NOT NULL DEFAULT 'PENDING',
                verification_status TEXT NOT NULL DEFAULT 'NOT_VERIFIED',
                verification_score INTEGER NOT NULL DEFAULT 0,
                verification_reason TEXT,
                uploaded_by_role TEXT,
                uploaded_by_name TEXT,
                uploaded_at TEXT NOT NULL,
                verified_at TEXT
            )
            """
        )

        # Safe migrations for existing local DBs.
        ensure_column(conn, "gap_tasks", "closed_at", "closed_at TEXT")
        ensure_column(conn, "gap_tasks", "verification_notes", "verification_notes TEXT")
        ensure_column(conn, "gap_tasks", "evidence_note", "evidence_note TEXT")

        ensure_column(conn, "task_evidence", "tender_id", "tender_id INTEGER NOT NULL DEFAULT 0")
        ensure_column(conn, "task_evidence", "original_filename", "original_filename TEXT NOT NULL DEFAULT ''")
        ensure_column(conn, "task_evidence", "stored_filename", "stored_filename TEXT NOT NULL DEFAULT ''")
        ensure_column(conn, "task_evidence", "file_path", "file_path TEXT NOT NULL DEFAULT ''")
        ensure_column(conn, "task_evidence", "mime_type", "mime_type TEXT")
        ensure_column(conn, "task_evidence", "file_size", "file_size INTEGER NOT NULL DEFAULT 0")
        ensure_column(conn, "task_evidence", "evidence_note", "evidence_note TEXT")
        ensure_column(conn, "task_evidence", "extracted_text", "extracted_text TEXT")
        ensure_column(conn, "task_evidence", "extraction_status", "extraction_status TEXT NOT NULL DEFAULT 'PENDING'")
        ensure_column(conn, "task_evidence", "verification_status", "verification_status TEXT NOT NULL DEFAULT 'NOT_VERIFIED'")
        ensure_column(conn, "task_evidence", "verification_score", "verification_score INTEGER NOT NULL DEFAULT 0")
        ensure_column(conn, "task_evidence", "verification_reason", "verification_reason TEXT")
        ensure_column(conn, "task_evidence", "uploaded_by_role", "uploaded_by_role TEXT")
        ensure_column(conn, "task_evidence", "uploaded_by_name", "uploaded_by_name TEXT")
        ensure_column(conn, "task_evidence", "uploaded_at", "uploaded_at TEXT")
        ensure_column(conn, "task_evidence", "verified_at", "verified_at TEXT")

        seed_team_members(conn)
        conn.commit()


def seed_team_members(conn: sqlite3.Connection) -> None:
    members = [
        ("أحمد - مدير المنافسة", "tender_manager", "مدير المنافسة", "إدارة العروض", "tender.manager@munjiz.local"),
        ("سارة - المهندس الفني", "technical_engineer", "مهندس فني", "الفريق الفني", "technical.engineer@munjiz.local"),
        ("خالد - مسؤول المشتريات", "procurement_officer", "مسؤول المشتريات", "المشتريات", "procurement@munjiz.local"),
        ("نورة - مسؤولة الجودة والامتثال", "quality_compliance", "مسؤول الجودة والامتثال", "الجودة والامتثال", "quality@munjiz.local"),
        ("ماجد - مسؤول المشاريع", "project_manager", "مسؤول المشاريع", "إدارة المشاريع", "projects@munjiz.local"),
        ("فيصل - المسؤول المالي", "finance_officer", "المسؤول المالي", "الإدارة المالية", "finance@munjiz.local"),
        ("ريم - مسؤولة المستندات", "document_controller", "مسؤول المستندات", "الوثائق والتسليم", "documents@munjiz.local"),
    ]

    for name, role, title, department, email in members:
        conn.execute(
            """
            INSERT OR IGNORE INTO team_members
            (name, role, title, department, email, active, created_at)
            VALUES (?, ?, ?, ?, ?, 1, ?)
            """,
            (name, role, title, department, email, now_iso()),
        )


def get_team_member_by_role(conn: sqlite3.Connection, role: str) -> dict[str, Any]:
    row = conn.execute(
        "SELECT * FROM team_members WHERE role = ? AND active = 1",
        (role,),
    ).fetchone()

    if row:
        return dict(row)

    fallback = conn.execute(
        "SELECT * FROM team_members WHERE role = 'tender_manager' AND active = 1"
    ).fetchone()

    return dict(fallback) if fallback else {
        "name": "مدير المنافسة",
        "role": "tender_manager",
        "title": "مدير المنافسة",
        "department": "إدارة العروض",
        "email": None,
    }


def normalize_status(value: str | None) -> str:
    return (value or "").strip().lower()


def is_requirement_covered(status: str | None) -> bool:
    normalized = normalize_status(status)
    if not normalized:
        return False

    negative_words = [
        "غير مغطى", "غير مغطاة", "غير مكتمل", "غير مكتملة",
        "not covered", "uncovered", "missing", "gap", "open", "pending",
        "needs evidence", "needs review", "بحاجة", "ناقص", "ناقصة", "مفتوح", "مفتوحة",
    ]

    partial_words = ["partial", "partially", "جزئي", "جزئية", "مغطى جزئيًا", "مغطى جزئياً"]

    covered_words = [
        "covered", "complete", "completed", "closed",
        "مغطى", "مغطاة", "مكتمل", "مكتملة", "مغلق", "مغلقة", "جاهز", "جاهزة",
    ]

    if any(word in normalized for word in negative_words):
        return False

    if any(word in normalized for word in partial_words):
        return False

    return any(word in normalized for word in covered_words)


def classify_gap(requirement_title: str, category: str | None = None) -> dict[str, Any]:
    text = f"{requirement_title} {category or ''}".lower()

    if any(word in text for word in ["iso", "quality", "compliance", "cyber", "security", "شهادة", "جودة", "امتثال", "أمن"]):
        return {
            "owner_role": "quality_compliance",
            "priority": "عالية",
            "evidence_type": "شهادة / سياسة امتثال",
            "impact_score": 10,
            "instruction": "ارفع شهادة سارية أو سياسة معتمدة أو خطاب رسمي يثبت الامتثال المطلوب.",
        }

    if any(word in text for word in ["authorization", "manufacturer", "vendor", "supplier", "distributor", "اعتماد", "تفويض", "مصنع", "مورد", "وكيل"]):
        return {
            "owner_role": "procurement_officer",
            "priority": "عالية",
            "evidence_type": "خطاب تفويض / اعتماد مورد",
            "impact_score": 10,
            "instruction": "ارفع خطاب تفويض من المصنع أو المورد أو ما يثبت توفر المنتج والدعم والضمان.",
        }

    if any(word in text for word in ["implementation", "methodology", "testing", "sat", "fat", "technical", "plan", "تنفيذ", "منهجية", "اختبار", "تشغيل", "فني"]):
        return {
            "owner_role": "technical_engineer",
            "priority": "متوسطة",
            "evidence_type": "خطة فنية / منهجية تنفيذ",
            "impact_score": 7,
            "instruction": "ارفع خطة تنفيذ أو اختبار وتشغيل أو مستند فني يثبت تغطية المتطلب.",
        }

    if any(word in text for word in ["experience", "project", "reference", "completion", "خبرة", "سابقة", "مشروع", "إنجاز", "اعمال", "أعمال"]):
        return {
            "owner_role": "project_manager",
            "priority": "عالية",
            "evidence_type": "سابقة أعمال / شهادة إنجاز",
            "impact_score": 9,
            "instruction": "ارفع شهادة إنجاز أو سابقة أعمال مشابهة تثبت قدرة الشركة على تنفيذ هذا المتطلب.",
        }

    if any(word in text for word in ["financial", "guarantee", "bank", "price", "ضمان", "مالي", "بنك", "تسعير", "سعر"]):
        return {
            "owner_role": "finance_officer",
            "priority": "عالية",
            "evidence_type": "مستند مالي / ضمان",
            "impact_score": 8,
            "instruction": "ارفع المستند المالي المطلوب مثل الضمان أو خطاب البنك أو ما يثبت الجاهزية المالية.",
        }

    if any(word in text for word in ["document", "attachment", "form", "ملف", "مستند", "مرفق", "نموذج"]):
        return {
            "owner_role": "document_controller",
            "priority": "متوسطة",
            "evidence_type": "مستند مطلوب",
            "impact_score": 6,
            "instruction": "ارفع المستند المطلوب أو النموذج المكتمل حسب كراسة المنافسة.",
        }

    return {
        "owner_role": "tender_manager",
        "priority": "متوسطة",
        "evidence_type": "دليل داعم",
        "impact_score": 5,
        "instruction": "راجع المتطلب وحدد الدليل المناسب لإغلاق الفجوة.",
    }


def fetch_tender_or_404(conn: sqlite3.Connection, tender_id: int) -> dict[str, Any]:
    row = conn.execute("SELECT * FROM tenders WHERE id = ?", (tender_id,)).fetchone()
    if not row:
        raise HTTPException(status_code=404, detail="Tender not found")
    return dict(row)


def fetch_task_or_404(conn: sqlite3.Connection, task_id: int) -> dict[str, Any]:
    row = conn.execute("SELECT * FROM gap_tasks WHERE id = ?", (task_id,)).fetchone()
    if not row:
        raise HTTPException(status_code=404, detail="Task not found")
    return dict(row)


def fetch_requirements(conn: sqlite3.Connection, tender_id: int) -> list[dict[str, Any]]:
    try:
        rows = conn.execute(
            "SELECT * FROM requirements WHERE tender_id = ? ORDER BY id ASC",
            (tender_id,),
        ).fetchall()
        return rows_to_list(rows)
    except sqlite3.OperationalError:
        return []


def create_task_event(conn: sqlite3.Connection, task_id: int, event_type: str, message: str) -> None:
    conn.execute(
        "INSERT INTO gap_task_events (task_id, event_type, message, created_at) VALUES (?, ?, ?, ?)",
        (task_id, event_type, message, now_iso()),
    )


def build_tasks_summary(tasks: list[dict[str, Any]]) -> dict[str, Any]:
    total = len(tasks)
    open_count = len([task for task in tasks if task.get("status") == "OPEN"])
    in_progress_count = len([task for task in tasks if task.get("status") == "IN_PROGRESS"])
    waiting_review_count = len([task for task in tasks if task.get("status") == "WAITING_REVIEW"])
    waiting_manager_count = len([task for task in tasks if task.get("status") == "WAITING_MANAGER_APPROVAL"])
    closed_count = len([task for task in tasks if task.get("status") == "CLOSED"])
    blocked_count = len([task for task in tasks if task.get("status") == "BLOCKED"])
    high_priority_open = len([task for task in tasks if task.get("priority") == "عالية" and task.get("status") not in ("CLOSED", "WAITING_MANAGER_APPROVAL")])

    total_impact = sum(int(task.get("impact_score") or 0) for task in tasks)
    closed_impact = sum(int(task.get("impact_score") or 0) for task in tasks if task.get("status") == "CLOSED")

    closure_score = 100 if total_impact == 0 else round((closed_impact / total_impact) * 100)

    if total == 0:
        decision = "لا توجد مهام فجوات"
        recommendation = "شغّل توليد مهام الفجوات بعد استخراج المتطلبات."
    elif high_priority_open > 0:
        decision = "دخول مشروط"
        recommendation = f"أغلق {high_priority_open} فجوة عالية التأثير قبل اعتماد التقديم."
    elif closure_score >= 85:
        decision = "جاهز للتقديم"
        recommendation = "الفجوات الحرجة مغلقة. انتقل إلى مراجعة القرار النهائي."
    else:
        decision = "قيد التجهيز"
        recommendation = "استكمل المهام المفتوحة ثم أعد التحقق."

    return {
        "total": total,
        "open": open_count,
        "in_progress": in_progress_count,
        "waiting_review": waiting_review_count,
        "waiting_manager": waiting_manager_count,
        "closed": closed_count,
        "blocked": blocked_count,
        "high_priority_open": high_priority_open,
        "closure_score": closure_score,
        "decision": decision,
        "recommendation": recommendation,
    }


def safe_filename(filename: str) -> str:
    clean_name = re.sub(r"[^A-Za-z0-9_.\-\u0600-\u06FF ]+", "_", filename).strip()
    return clean_name or "evidence_file"


def extract_text_from_file(path: Path, mime_type: str | None) -> tuple[str, str]:
    suffix = path.suffix.lower()
    mime = (mime_type or "").lower()

    try:
        if suffix in [".txt", ".md", ".csv"] or mime.startswith("text/"):
            return path.read_text(encoding="utf-8", errors="ignore"), "EXTRACTED"

        if suffix == ".pdf" or mime == "application/pdf":
            try:
                import pypdf  # type: ignore
                reader = pypdf.PdfReader(str(path))
                text_parts = [(page.extract_text() or "") for page in reader.pages]
                extracted = "\n".join(text_parts).strip()
                return (extracted or "تم حفظ ملف PDF وربطه بالمهمة.", "EXTRACTED" if extracted else "SAVED_ONLY")
            except Exception:
                return "تم حفظ ملف PDF وربطه بالمهمة.", "SAVED_ONLY"

        if suffix == ".docx" or "wordprocessingml" in mime:
            try:
                import docx  # type: ignore
                document = docx.Document(str(path))
                extracted = "\n".join([p.text for p in document.paragraphs]).strip()
                return (extracted or "تم حفظ ملف DOCX وربطه بالمهمة.", "EXTRACTED" if extracted else "SAVED_ONLY")
            except Exception:
                return "تم حفظ ملف DOCX وربطه بالمهمة.", "SAVED_ONLY"

        return "تم حفظ الملف وربطه بالمهمة.", "SAVED_ONLY"
    except Exception:
        return "تم حفظ الملف وربطه بالمهمة.", "SAVED_ONLY"


def tokenize_for_matching(text: str) -> set[str]:
    words = re.findall(r"[\w\u0600-\u06FF]+", text.lower())
    stop_words = {
        "في", "من", "على", "إلى", "الى", "عن", "مع", "هذا", "هذه", "ذلك",
        "the", "and", "or", "for", "with", "to", "of", "a", "an", "is", "are",
        "must", "shall",
        "لدينا", "القدرة", "سيتم", "تقديم", "الخدمات", "حسب", "الحاجة", "شكرا", "شكراً", "نؤكد",
        "جيدا", "جيد", "مطلوب", "مطلوبة", "بشكل", "المطلوبة", "نحن", "سوف", "نقوم", "تنفيذ", "المشروع"
    }
    return {word for word in words if len(word) >= 3 and word not in stop_words}


def verify_evidence_against_task(task: dict[str, Any], evidence: dict[str, Any]) -> dict[str, Any]:
    extracted_text = evidence.get("extracted_text") or ""
    evidence_note = evidence.get("evidence_note") or ""
    task_text = " ".join([
        str(task.get("requirement_title") or ""),
        str(task.get("evidence_type") or ""),
        str(task.get("evidence_instruction") or ""),
    ])

    combined_text = f"{extracted_text}\n{evidence_note}".strip()
    matched_terms = sorted(list(tokenize_for_matching(task_text).intersection(tokenize_for_matching(combined_text))))

    score = 0
    reasons: list[str] = []

    has_file = int(evidence.get("file_size") or 0) > 0
    
    if has_file:
        score += 20
        reasons.append("يوجد ملف مرفق.")

    if len(combined_text) >= 50:
        score += 20
        reasons.append("محتوى الدليل ذو طول كافٍ للمراجعة.")
    elif len(combined_text) > 0:
        score += 5
        reasons.append("محتوى الدليل قصير جداً.")

    if matched_terms:
        score += min(60, len(matched_terms) * 15)
        reasons.append(f"توجد مؤشرات مطابقة للمتطلب: {', '.join(matched_terms[:6])}.")

    missing_items: list[str] = []
    if not has_file:
        missing_items.append("ملف دليل فعلي")
    if len(combined_text) < 50:
        missing_items.append("محتوى تفصيلي كافٍ للمراجعة")
    if len(matched_terms) < 3:
        missing_items.append("مؤشرات فنية أو كلمات مفتاحية مطابقة لمتطلبات المهمة")

    # Hard rules to prevent false acceptances of weak generic evidence
    if len(matched_terms) < 3:
        score = min(score, 55)
    if len(combined_text) < 50:
        score = min(score, 55)

    if score >= 75:
        recommended_action = "يمكن إغلاق هذه الفجوة. تأكد من حفظ الدليل في سجل التسليم النهائي."
        decision = "مقبول"
    elif score >= 40:
        recommended_action = "الدليل جزئي أو عام. يُنصح برفع دليل فني مفصل يوضح الامتثال الدقيق للمتطلب."
        decision = "يستلزم مراجعة"
    else:
        recommended_action = "الدليل غير كافٍ. ارفع ملفًا فعليًا يتضمن البيانات المطلوبة قبل إعادة التحقق."
        decision = "مرفوض"

    return {
        "verification_status": "ACCEPTED" if score >= 75 else "REJECTED",
        "task_status": "CLOSED" if score >= 75 else "WAITING_REVIEW",
        "verification_score": score,
        "verification_reason": ("تم قبول الدليل وإغلاق الفجوة. " if score >= 75 else "تم رفض الدليل لعدم اكتماله. ") + " ".join(reasons),
        "decision": decision,
        "matched_indicators": matched_terms[:6],
        "missing_items": missing_items,
        "reasoning": reasons,
        "recommended_action": recommended_action,
        "manager_approval_required": False,
    }

class ManagerRejectRequest(BaseModel):
    rejection_note: str



class TaskUpdateRequest(BaseModel):
    status: Optional[str] = None
    owner_role: Optional[str] = None
    priority: Optional[str] = None
    evidence_note: Optional[str] = None
    verification_notes: Optional[str] = None


class VerifyTaskRequest(BaseModel):
    evidence_note: Optional[str] = Field(default=None)


@router.on_event("startup")
def startup_gap_closure() -> None:
    ensure_gap_tables()


@router.get("/team")
def list_team_members() -> dict[str, Any]:
    ensure_gap_tables()
    with get_connection() as conn:
        rows = conn.execute("SELECT * FROM team_members WHERE active = 1 ORDER BY id ASC").fetchall()
    return {"team": rows_to_list(rows)}


@router.get("/tasks")
def list_gap_tasks(
    tender_id: Optional[int] = Query(default=None),
    owner_role: Optional[str] = Query(default=None),
    status: Optional[str] = Query(default=None),
) -> dict[str, Any]:
    ensure_gap_tables()

    conditions: list[str] = []
    params: list[Any] = []

    if tender_id is not None:
        conditions.append("tender_id = ?")
        params.append(tender_id)

    if owner_role:
        conditions.append("owner_role = ?")
        params.append(owner_role)

    if status:
        conditions.append("status = ?")
        params.append(status)

    where_clause = "WHERE " + " AND ".join(conditions) if conditions else ""

    with get_connection() as conn:
        rows = conn.execute(
            f"""
            SELECT *
            FROM gap_tasks
            {where_clause}
            ORDER BY
                CASE status
                    WHEN 'OPEN' THEN 1
                    WHEN 'IN_PROGRESS' THEN 2
                    WHEN 'WAITING_REVIEW' THEN 3
                    WHEN 'CLOSED' THEN 4
                    ELSE 5
                END,
                CASE priority
                    WHEN 'عالية' THEN 1
                    WHEN 'متوسطة' THEN 2
                    ELSE 3
                END,
                id DESC
            """,
            params,
        ).fetchall()

    return {"tasks": rows_to_list(rows)}


@router.get("/workbench/{owner_role}")
def employee_workbench(owner_role: str) -> dict[str, Any]:
    ensure_gap_tables()

    with get_connection() as conn:
        member = get_team_member_by_role(conn, owner_role)
        rows = conn.execute(
            """
            SELECT *
            FROM gap_tasks
            WHERE owner_role = ?
            ORDER BY
                CASE status
                    WHEN 'OPEN' THEN 1
                    WHEN 'IN_PROGRESS' THEN 2
                    WHEN 'WAITING_REVIEW' THEN 3
                    WHEN 'CLOSED' THEN 4
                    ELSE 5
                END,
                CASE priority
                    WHEN 'عالية' THEN 1
                    WHEN 'متوسطة' THEN 2
                    ELSE 3
                END,
                id DESC
            """,
            (owner_role,),
        ).fetchall()
        tasks = rows_to_list(rows)

    return {"member": member, "tasks": tasks, "summary": build_tasks_summary(tasks)}


@router.post("/tenders/{tender_id}/generate")
def generate_gap_tasks(tender_id: int) -> dict[str, Any]:
    ensure_gap_tables()

    created_tasks: list[dict[str, Any]] = []
    skipped_tasks: list[dict[str, Any]] = []

    with get_connection() as conn:
        tender = fetch_tender_or_404(conn, tender_id)
        requirements = fetch_requirements(conn, tender_id)

        if not requirements:
            raise HTTPException(
                status_code=400,
                detail="No requirements found for this tender. Extract or create requirements first.",
            )

        due_date = (datetime.utcnow() + timedelta(days=3)).date().isoformat()

        for requirement in requirements:
            status = requirement.get("status")
            requirement_id = requirement.get("id")
            title = requirement.get("title") or "متطلب بدون عنوان"
            category = requirement.get("category") or "عام"

            if is_requirement_covered(status):
                skipped_tasks.append({"requirement_id": requirement_id, "reason": "Requirement already covered"})
                continue

            existing = conn.execute(
                """
                SELECT *
                FROM gap_tasks
                WHERE tender_id = ?
                  AND requirement_id = ?
                  AND status != 'CLOSED'
                """,
                (tender_id, requirement_id),
            ).fetchone()

            if existing:
                skipped_tasks.append({"requirement_id": requirement_id, "reason": "Open task already exists"})
                continue

            classification = classify_gap(title, category)
            member = get_team_member_by_role(conn, classification["owner_role"])

            task_title = f"طلب دليل: {title[:90]}"
            task_description = (
                f"تم اكتشاف فجوة في المنافسة: {tender.get('title', '')}. "
                f"المتطلب يحتاج دليلًا قبل اعتماد قرار التقديم. "
                f"التصنيف: {category}. "
                f"حالة المتطلب الحالية: {status or 'غير محددة'}. "
                f"الإجراء المطلوب: {classification['instruction']}"
            )

            cursor = conn.execute(
                """
                INSERT INTO gap_tasks (
                    tender_id, requirement_id, requirement_title, title, description,
                    owner_role, owner_name, owner_department, priority, status,
                    evidence_type, evidence_instruction, impact_score, due_date,
                    verification_status, created_by, created_at, updated_at
                )
                VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, 'OPEN', ?, ?, ?, ?, 'NOT_SUBMITTED', 'Gap Assignment Agent', ?, ?)
                """,
                (
                    tender_id, requirement_id, title, task_title, task_description,
                    member["role"], member["name"], member["department"],
                    classification["priority"], classification["evidence_type"],
                    classification["instruction"], classification["impact_score"], due_date,
                    now_iso(), now_iso(),
                ),
            )

            task_id = int(cursor.lastrowid)
            create_task_event(conn, task_id, "TASK_CREATED", f"تم إنشاء مهمة فجوة وإسنادها إلى {member['title']}.")
            created_row = conn.execute("SELECT * FROM gap_tasks WHERE id = ?", (task_id,)).fetchone()
            if created_row:
                created_tasks.append(dict(created_row))

        conn.commit()

    return {
        "message": "Gap tasks generated successfully",
        "tender_id": tender_id,
        "created_count": len(created_tasks),
        "skipped_count": len(skipped_tasks),
        "created_tasks": created_tasks,
        "skipped_tasks": skipped_tasks,
    }


@router.post("/tenders/{tender_id}/reset-scenario")
def reset_scenario(tender_id: int) -> dict[str, Any]:
    ensure_gap_tables()
    
    if tender_id != 4:
        raise HTTPException(status_code=403, detail="Scenario reset is only allowed for the scenario tender.")
        
    with get_connection() as conn:
        tender = fetch_tender_or_404(conn, tender_id)
        
        # Fetch all tasks ordered by priority and ID
        rows = conn.execute(
            """
            SELECT * FROM gap_tasks 
            WHERE tender_id = ? 
            ORDER BY 
                CASE priority 
                    WHEN 'عالية' THEN 1 
                    WHEN 'متوسطة' THEN 2 
                    ELSE 3 
                END, 
                id ASC
            """,
            (tender_id,)
        ).fetchall()
        
        tasks = rows_to_list(rows)
        if not tasks:
            return {"message": "No tasks to reset", "reset_count": 0}
            
        high_tasks = [t for t in tasks if t["priority"] == "عالية"]
        med_tasks = [t for t in tasks if t["priority"] == "متوسطة"]
        low_tasks = [t for t in tasks if t["priority"] not in ("عالية", "متوسطة")]
        
        def assign_status(task_list: list[dict], pattern: list[str]) -> None:
            for i, t in enumerate(task_list):
                if i < len(pattern):
                    t["_new_status"] = pattern[i]
                else:
                    t["_new_status"] = pattern[-1]
                    
        # High priority tasks (عالية): 1 -> WAITING_MANAGER_APPROVAL, 1 -> CLOSED, rest -> OPEN
        assign_status(high_tasks, ["WAITING_MANAGER_APPROVAL", "CLOSED", "OPEN"])
        
        # Medium priority tasks (متوسطة): 1 -> WAITING_REVIEW, 1 -> CLOSED, rest -> IN_PROGRESS
        assign_status(med_tasks, ["WAITING_REVIEW", "CLOSED", "IN_PROGRESS"])
        
        # Low priority tasks: all -> OPEN
        assign_status(low_tasks, ["OPEN"])
        
        all_reset_tasks = high_tasks + med_tasks + low_tasks
        
        base_time = datetime.utcnow() - timedelta(days=3)
        
        for t in all_reset_tasks:
            status = t["_new_status"]
            task_id = t["id"]
            
            # Clear existing events
            conn.execute("DELETE FROM gap_task_events WHERE task_id = ?", (task_id,))
            
            # Define verif_status and closed_at
            verif_status = "NOT_SUBMITTED"
            closed_at = None
            
            if status == "CLOSED":
                verif_status = "ACCEPTED"
                closed_at = now_iso()
            elif status == "WAITING_MANAGER_APPROVAL":
                verif_status = "ACCEPTED"
            elif status == "WAITING_REVIEW":
                verif_status = "PENDING"
            elif status in ("IN_PROGRESS", "OPEN"):
                verif_status = "NOT_SUBMITTED"
                
            conn.execute(
                """
                UPDATE gap_tasks 
                SET status = ?, verification_status = ?, verification_notes = NULL, closed_at = ?, updated_at = ?
                WHERE id = ?
                """,
                (status, verif_status, closed_at, now_iso(), task_id)
            )
            
            # Add synthetic events
            create_task_event(conn, task_id, "TASK_CREATED", f"تم إنشاء مهمة فجوة وإسنادها إلى {t.get('owner_role')}.")
            
            if status in ("IN_PROGRESS", "WAITING_REVIEW", "WAITING_MANAGER_APPROVAL", "CLOSED"):
                event_time = (base_time + timedelta(days=1)).isoformat() + "Z"
                conn.execute(
                    "INSERT INTO gap_task_events (task_id, event_type, message, created_at) VALUES (?, ?, ?, ?)",
                    (task_id, "STATUS_CHANGED", "تم بدء العمل على المهمة.", event_time)
                )
            
            if status in ("WAITING_REVIEW", "WAITING_MANAGER_APPROVAL", "CLOSED"):
                event_time = (base_time + timedelta(days=2)).isoformat() + "Z"
                conn.execute(
                    "INSERT INTO gap_task_events (task_id, event_type, message, created_at) VALUES (?, ?, ?, ?)",
                    (task_id, "EVIDENCE_UPLOADED", "تم رفع دليل للمراجعة.", event_time)
                )
                
            if status in ("WAITING_MANAGER_APPROVAL", "CLOSED"):
                event_time = (base_time + timedelta(hours=50)).isoformat() + "Z"
                conn.execute(
                    "INSERT INTO gap_task_events (task_id, event_type, message, created_at) VALUES (?, ?, ?, ?)",
                    (task_id, "EVIDENCE_VERIFIED", "تم التحقق من الدليل بنجاح.", event_time)
                )
                
            if status == "CLOSED":
                event_time = (base_time + timedelta(hours=55)).isoformat() + "Z"
                conn.execute(
                    "INSERT INTO gap_task_events (task_id, event_type, message, created_at) VALUES (?, ?, ?, ?)",
                    (task_id, "MANAGER_APPROVED", "تم اعتماد الدليل وإغلاق الفجوة.", event_time)
                )
                
            # Update evidence status safely
            if status in ("WAITING_MANAGER_APPROVAL", "CLOSED"):
                conn.execute("UPDATE task_evidence SET verification_status = 'ACCEPTED' WHERE task_id = ?", (task_id,))
            else:
                conn.execute("UPDATE task_evidence SET verification_status = 'NOT_VERIFIED' WHERE task_id = ?", (task_id,))
                
        conn.commit()
        
        # Build summary
        updated_rows = conn.execute("SELECT * FROM gap_tasks WHERE tender_id = ?", (tender_id,)).fetchall()
        summary = build_tasks_summary(rows_to_list(updated_rows))
        
    return {
        "message": "تمت إعادة تهيئة سيناريو العرض بنجاح",
        "reset_count": len(all_reset_tasks),
        "summary": summary
    }


@router.patch("/tasks/{task_id}")
def update_gap_task(task_id: int, payload: TaskUpdateRequest) -> dict[str, Any]:
    ensure_gap_tables()
    allowed_statuses = {"OPEN", "IN_PROGRESS", "WAITING_REVIEW", "WAITING_MANAGER_APPROVAL", "CLOSED", "BLOCKED"}

    with get_connection() as conn:
        task = fetch_task_or_404(conn, task_id)

        updates: list[str] = []
        params: list[Any] = []
        event_messages: list[str] = []

        if payload.status is not None:
            if payload.status not in allowed_statuses:
                raise HTTPException(status_code=400, detail="Invalid task status")
            updates.append("status = ?")
            params.append(payload.status)
            if payload.status == "CLOSED":
                updates.append("closed_at = ?")
                params.append(now_iso())
            event_messages.append(f"تم تحديث حالة المهمة إلى {payload.status}.")

        if payload.owner_role is not None:
            member = get_team_member_by_role(conn, payload.owner_role)
            updates.extend(["owner_role = ?", "owner_name = ?", "owner_department = ?"])
            params.extend([member["role"], member["name"], member["department"]])
            event_messages.append(f"تم تغيير مالك المهمة إلى {member['title']}.")

        if payload.priority is not None:
            updates.append("priority = ?")
            params.append(payload.priority)
            event_messages.append(f"تم تعديل أولوية المهمة إلى {payload.priority}.")

        if payload.evidence_note is not None:
            updates.append("evidence_note = ?")
            params.append(payload.evidence_note)
            updates.append("verification_status = ?")
            params.append("SUBMITTED")
            event_messages.append("تم تسجيل ملاحظة دليل للمهمة.")

        if payload.verification_notes is not None:
            updates.append("verification_notes = ?")
            params.append(payload.verification_notes)
            event_messages.append("تم تحديث ملاحظات التحقق.")

        if not updates:
            raise HTTPException(status_code=400, detail="No fields to update")

        updates.append("updated_at = ?")
        params.append(now_iso())
        params.append(task_id)

        conn.execute(f"UPDATE gap_tasks SET {', '.join(updates)} WHERE id = ?", params)

        for message in event_messages:
            create_task_event(conn, task_id, "TASK_UPDATED", message)

        conn.commit()
        updated = conn.execute("SELECT * FROM gap_tasks WHERE id = ?", (task_id,)).fetchone()

    return {"message": "Task updated", "task": row_to_dict(updated)}


@router.post("/tasks/{task_id}/evidence/upload")
async def upload_task_evidence(
    task_id: int,
    file: UploadFile = File(...),
    evidence_note: Optional[str] = Form(default=None),
    uploaded_by_role: Optional[str] = Form(default=None),
) -> dict[str, Any]:
    ensure_gap_tables()

    with get_connection() as conn:
        task = fetch_task_or_404(conn, task_id)

        if task.get("status") == "CLOSED":
            raise HTTPException(status_code=400, detail="Task is already closed")

        member = get_team_member_by_role(conn, uploaded_by_role or str(task.get("owner_role") or "tender_manager"))

    original_filename = safe_filename(file.filename or "evidence_file")
    extension = Path(original_filename).suffix
    stored_filename = f"task_{task_id}_{uuid.uuid4().hex}{extension}"
    file_path = get_evidence_storage_dir() / stored_filename

    size = 0
    with file_path.open("wb") as buffer:
        while True:
            chunk = await file.read(1024 * 1024)
            if not chunk:
                break
            size += len(chunk)
            buffer.write(chunk)

    if size <= 0:
        raise HTTPException(status_code=400, detail="Uploaded file is empty")

    extracted_text, extraction_status = extract_text_from_file(file_path, file.content_type)

    with get_connection() as conn:
        task = fetch_task_or_404(conn, task_id)

        cursor = conn.execute(
            """
            INSERT INTO task_evidence (
                task_id, tender_id, original_filename, stored_filename, file_path,
                mime_type, file_size, evidence_note, extracted_text, extraction_status,
                verification_status, verification_score, uploaded_by_role, uploaded_by_name, uploaded_at
            )
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, 'NOT_VERIFIED', 0, ?, ?, ?)
            """,
            (
                task_id, int(task["tender_id"]), original_filename, stored_filename, str(file_path),
                file.content_type, size, evidence_note, extracted_text, extraction_status,
                member["role"], member["name"], now_iso(),
            ),
        )

        evidence_id = int(cursor.lastrowid)

        conn.execute(
            """
            UPDATE gap_tasks
            SET
                status = 'WAITING_REVIEW',
                verification_status = 'SUBMITTED',
                evidence_note = ?,
                verification_notes = ?,
                updated_at = ?
            WHERE id = ?
            """,
            (
                evidence_note or f"تم رفع ملف دليل: {original_filename}",
                f"تم استلام الدليل وربطه بسجل المهمة: {original_filename}.",
                now_iso(),
                task_id,
            ),
        )

        create_task_event(conn, task_id, "EVIDENCE_UPLOADED", f"تم استلام دليل فعلي وربطه بالمهمة: {original_filename}.")
        conn.commit()

        evidence = conn.execute("SELECT * FROM task_evidence WHERE id = ?", (evidence_id,)).fetchone()
        updated_task = conn.execute("SELECT * FROM gap_tasks WHERE id = ?", (task_id,)).fetchone()

    return {
        "message": "تم استلام الدليل وربطه بالمهمة بنجاح.",
        "task": row_to_dict(updated_task),
        "evidence": row_to_dict(evidence),
    }


@router.get("/tasks/{task_id}/evidence")
def list_task_evidence(task_id: int) -> dict[str, Any]:
    ensure_gap_tables()

    with get_connection() as conn:
        fetch_task_or_404(conn, task_id)
        rows = conn.execute(
            "SELECT * FROM task_evidence WHERE task_id = ? ORDER BY id DESC",
            (task_id,),
        ).fetchall()

    return {"task_id": task_id, "evidence": rows_to_list(rows)}


@router.get("/evidence/{evidence_id}/download")
def download_evidence(evidence_id: int) -> FileResponse:
    ensure_gap_tables()

    with get_connection() as conn:
        row = conn.execute("SELECT * FROM task_evidence WHERE id = ?", (evidence_id,)).fetchone()
        if not row:
            raise HTTPException(status_code=404, detail="Evidence not found")
        evidence = dict(row)

    file_path = Path(evidence["file_path"])
    if not file_path.exists():
        raise HTTPException(status_code=404, detail="Evidence file not found on disk")

    return FileResponse(
        path=str(file_path),
        filename=evidence.get("original_filename") or file_path.name,
        media_type=evidence.get("mime_type") or "application/octet-stream",
    )


@router.post("/tasks/{task_id}/verify")
def verify_gap_task(task_id: int, payload: VerifyTaskRequest | None = None) -> dict[str, Any]:
    ensure_gap_tables()

    with get_connection() as conn:
        task = fetch_task_or_404(conn, task_id)
        evidence_row = conn.execute(
            "SELECT * FROM task_evidence WHERE task_id = ? ORDER BY id DESC LIMIT 1",
            (task_id,),
        ).fetchone()

        if evidence_row:
            evidence = dict(evidence_row)
            result = verify_evidence_against_task(task, evidence)
            
            manager_approval_required = False
            if result["task_status"] == "CLOSED" and task.get("priority") == "عالية":
                result["task_status"] = "WAITING_MANAGER_APPROVAL"
                result["verification_status"] = "MANAGER_REVIEW"
                result["verification_reason"] = result["verification_reason"].replace("وإغلاق الفجوة", "وبانتظار اعتماد مدير المنافسة")
                manager_approval_required = True
                
            result["manager_approval_required"] = manager_approval_required

            closed_at = now_iso() if result["task_status"] == "CLOSED" else None


            conn.execute(
                """
                UPDATE task_evidence
                SET verification_status = ?, verification_score = ?, verification_reason = ?, verified_at = ?
                WHERE id = ?
                """,
                (
                    result["verification_status"],
                    result["verification_score"],
                    result["verification_reason"],
                    now_iso(),
                    evidence["id"],
                ),
            )

            conn.execute(
                """
                UPDATE gap_tasks
                SET verification_status = ?, verification_notes = ?, status = ?, closed_at = ?, updated_at = ?
                WHERE id = ?
                """,
                (
                    result["verification_status"],
                    result["verification_reason"],
                    result["task_status"],
                    closed_at,
                    now_iso(),
                    task_id,
                ),
            )

            create_task_event(conn, task_id, "TASK_VERIFIED", result["verification_reason"])
            conn.commit()

            updated_task = conn.execute("SELECT * FROM gap_tasks WHERE id = ?", (task_id,)).fetchone()
            updated_evidence = conn.execute("SELECT * FROM task_evidence WHERE id = ?", (evidence["id"],)).fetchone()

            verification_detail = {
                "decision": result.get("decision", ""),
                "confidence": result.get("verification_score", 0),
                "matched_indicators": result.get("matched_indicators", []),
                "missing_items": result.get("missing_items", []),
                "reasoning": result.get("reasoning", []),
                "recommended_action": result.get("recommended_action", ""),
                "manager_approval_required": result.get("manager_approval_required", False),
            }

            return {
                "message": result["verification_reason"],
                "task": row_to_dict(updated_task),
                "evidence": row_to_dict(updated_evidence),
                "verification_detail": verification_detail,
            }

        evidence_note = payload.evidence_note.strip() if payload and payload.evidence_note else ""

        manager_approval_required = False
        verification_detail = None
        if len(evidence_note) >= 15:
            if task.get("priority") == "عالية":
                verification_status = "MANAGER_REVIEW"
                status = "WAITING_MANAGER_APPROVAL"
                verification_notes = "تم قبول الدليل النصي وبانتظار اعتماد مدير المنافسة."
                closed_at = None
                event_message = "تم قبول الدليل النصي وبانتظار اعتماد مدير المنافسة."
                manager_approval_required = True
                decision = "يستلزم مراجعة"
            else:
                verification_status = "ACCEPTED"
                status = "CLOSED"
                verification_notes = "تم قبول الدليل النصي وإغلاق الفجوة."
                closed_at = now_iso()
                event_message = "تم قبول الدليل النصي وإغلاق الفجوة."
                decision = "مقبول"
                
            verification_detail = {
                "decision": decision,
                "confidence": 80,
                "matched_indicators": [],
                "missing_items": [],
                "reasoning": ["تم الاعتماد على ملاحظة الدليل النصية بدون ملف مرفق."],
                "recommended_action": "لا إجراء إضافي مطلوب، ولكن يُفضل إرفاق ملف لتوثيق أقوى.",
                "manager_approval_required": manager_approval_required,
            }
        else:
            verification_status = "REJECTED"
            status = "WAITING_REVIEW"
            verification_notes = "لا يوجد دليل مرفوع لهذه المهمة. ارفع ملفًا فعليًا ثم أعد التحقق."
            closed_at = None
            event_message = "تم رفض التحقق لعدم وجود دليل مرفوع."
            verification_detail = {
                "decision": "مرفوض",
                "confidence": 0,
                "matched_indicators": [],
                "missing_items": ["ملف دليل فعلي", "ملاحظة توضيحية"],
                "reasoning": ["لا يوجد محتوى كافٍ للتحقق."],
                "recommended_action": "ارفع ملفًا فعليًا يتضمن البيانات المطلوبة قبل إعادة التحقق.",
                "manager_approval_required": False,
            }

        conn.execute(
            """
            UPDATE gap_tasks
            SET evidence_note = ?, verification_status = ?, verification_notes = ?, status = ?, closed_at = ?, updated_at = ?
            WHERE id = ?
            """,
            (evidence_note, verification_status, verification_notes, status, closed_at, now_iso(), task_id),
        )

        create_task_event(conn, task_id, "TASK_VERIFIED", event_message)
        conn.commit()
        updated = conn.execute("SELECT * FROM gap_tasks WHERE id = ?", (task_id,)).fetchone()

    return {
        "message": event_message, 
        "task": row_to_dict(updated),
        "verification_detail": verification_detail,
    }


def build_task_timeline(conn: sqlite3.Connection, task_id: int) -> list[dict[str, Any]]:
    events = conn.execute(
        """
        SELECT
            id,
            'event' AS source_type,
            event_type AS type,
            message AS title,
            NULL AS filename,
            NULL AS evidence_id,
            NULL AS verification_status,
            NULL AS verification_score,
            created_at AS timestamp
        FROM gap_task_events
        WHERE task_id = ?
        """,
        (task_id,),
    ).fetchall()

    evidence_events = conn.execute(
        """
        SELECT
            id,
            'evidence' AS source_type,
            'EVIDENCE_RECORD' AS type,
            COALESCE(verification_reason, 'تم تسجيل دليل مرتبط بالمهمة.') AS title,
            original_filename AS filename,
            id AS evidence_id,
            verification_status,
            verification_score,
            uploaded_at AS timestamp
        FROM task_evidence
        WHERE task_id = ?
        """,
        (task_id,),
    ).fetchall()

    combined = rows_to_list(events) + rows_to_list(evidence_events)
    combined.sort(key=lambda item: item.get("timestamp") or "", reverse=True)
    return combined


@router.get("/tasks/{task_id}/timeline")
def task_timeline(task_id: int) -> dict[str, Any]:
    ensure_gap_tables()

    with get_connection() as conn:
        task = fetch_task_or_404(conn, task_id)
        timeline = build_task_timeline(conn, task_id)
        evidence_rows = conn.execute(
            "SELECT * FROM task_evidence WHERE task_id = ? ORDER BY id DESC",
            (task_id,),
        ).fetchall()

    return {"task": task, "evidence": rows_to_list(evidence_rows), "timeline": timeline}


@router.get("/tenders/{tender_id}/dashboard")
def tender_gap_dashboard(tender_id: int) -> dict[str, Any]:
    ensure_gap_tables()

    with get_connection() as conn:
        tender = fetch_tender_or_404(conn, tender_id)

        task_rows = conn.execute(
            "SELECT * FROM gap_tasks WHERE tender_id = ? ORDER BY id DESC",
            (tender_id,),
        ).fetchall()
        tasks = rows_to_list(task_rows)

        owner_rows = conn.execute(
            """
            SELECT
                owner_role,
                owner_name,
                owner_department,
                COUNT(*) AS total,
                SUM(CASE WHEN status = 'CLOSED' THEN 1 ELSE 0 END) AS closed,
                SUM(CASE WHEN status != 'CLOSED' THEN 1 ELSE 0 END) AS open_count,
                SUM(CASE WHEN priority = 'عالية' AND status != 'CLOSED' THEN 1 ELSE 0 END) AS high_open
            FROM gap_tasks
            WHERE tender_id = ?
            GROUP BY owner_role, owner_name, owner_department
            ORDER BY open_count DESC, total DESC
            """,
            (tender_id,),
        ).fetchall()

        evidence_rows = conn.execute(
            "SELECT * FROM task_evidence WHERE tender_id = ? ORDER BY uploaded_at DESC LIMIT 20",
            (tender_id,),
        ).fetchall()

        recent_event_rows = conn.execute(
            """
            SELECT
                e.id,
                e.task_id,
                t.title AS task_title,
                t.owner_name,
                e.event_type,
                e.message,
                e.created_at
            FROM gap_task_events e
            JOIN gap_tasks t ON t.id = e.task_id
            WHERE t.tender_id = ?
            ORDER BY e.id DESC
            LIMIT 30
            """,
            (tender_id,),
        ).fetchall()

    return {
        "tender": tender,
        "summary": build_tasks_summary(tasks),
        "owners": rows_to_list(owner_rows),
        "recent_evidence": rows_to_list(evidence_rows),
        "recent_events": rows_to_list(recent_event_rows),
    }


@router.get("/tasks/{task_id}/events")
def list_task_events(task_id: int) -> dict[str, Any]:
    ensure_gap_tables()

    with get_connection() as conn:
        rows = conn.execute(
            "SELECT * FROM gap_task_events WHERE task_id = ? ORDER BY id ASC",
            (task_id,),
        ).fetchall()

    return {"events": rows_to_list(rows)}


@router.get("/tenders/{tender_id}/summary")
def tender_gap_summary(tender_id: int) -> dict[str, Any]:
    ensure_gap_tables()

    with get_connection() as conn:
        fetch_tender_or_404(conn, tender_id)
        rows = conn.execute("SELECT * FROM gap_tasks WHERE tender_id = ?", (tender_id,)).fetchall()

    return {"tender_id": tender_id, "summary": build_tasks_summary(rows_to_list(rows))}


@router.delete("/tenders/{tender_id}/tasks")
def delete_tender_gap_tasks(tender_id: int) -> dict[str, Any]:
    ensure_gap_tables()

    with get_connection() as conn:
        fetch_tender_or_404(conn, tender_id)

        task_rows = conn.execute("SELECT id FROM gap_tasks WHERE tender_id = ?", (tender_id,)).fetchall()
        task_ids = [int(row["id"]) for row in task_rows]

        if task_ids:
            placeholders = ",".join(["?"] * len(task_ids))
            conn.execute(f"DELETE FROM gap_task_events WHERE task_id IN ({placeholders})", task_ids)
            conn.execute(f"DELETE FROM task_evidence WHERE task_id IN ({placeholders})", task_ids)

        conn.execute("DELETE FROM gap_tasks WHERE tender_id = ?", (tender_id,))
        conn.commit()

    return {"message": "Tender gap tasks deleted", "tender_id": tender_id, "deleted_count": len(task_ids)}

@router.get("/tenders/{tender_id}/manager-review")
def get_manager_review_tasks(tender_id: int) -> dict[str, Any]:
    ensure_gap_tables()
    with get_connection() as conn:
        fetch_tender_or_404(conn, tender_id)
        rows = conn.execute(
            """
            SELECT * FROM gap_tasks
            WHERE tender_id = ? AND status = 'WAITING_MANAGER_APPROVAL'
            ORDER BY priority ASC, id DESC
            """,
            (tender_id,),
        ).fetchall()
    return {"tasks": rows_to_list(rows)}

@router.post("/tasks/{task_id}/manager-approve")
def manager_approve_task(task_id: int) -> dict[str, Any]:
    ensure_gap_tables()
    with get_connection() as conn:
        task = fetch_task_or_404(conn, task_id)
        if task.get("status") != "WAITING_MANAGER_APPROVAL":
            raise HTTPException(status_code=400, detail="Task is not waiting for manager approval")
        
        conn.execute(
            """
            UPDATE gap_tasks
            SET status = 'CLOSED', verification_status = 'ACCEPTED', closed_at = ?, updated_at = ?
            WHERE id = ?
            """,
            (now_iso(), now_iso(), task_id),
        )
        create_task_event(conn, task_id, "MANAGER_APPROVED", "تم اعتماد الدليل وإغلاق الفجوة من قبل مدير المنافسة.")
        conn.commit()
        updated = conn.execute("SELECT * FROM gap_tasks WHERE id = ?", (task_id,)).fetchone()
    return {"message": "تم اعتماد الدليل وإغلاق الفجوة", "task": row_to_dict(updated)}

@router.post("/tasks/{task_id}/manager-reject")
def manager_reject_task(task_id: int, payload: ManagerRejectRequest) -> dict[str, Any]:
    ensure_gap_tables()
    with get_connection() as conn:
        task = fetch_task_or_404(conn, task_id)
        if task.get("status") != "WAITING_MANAGER_APPROVAL":
            raise HTTPException(status_code=400, detail="Task is not waiting for manager approval")
        
        conn.execute(
            """
            UPDATE gap_tasks
            SET status = 'IN_PROGRESS', verification_status = 'REJECTED', verification_notes = ?, updated_at = ?
            WHERE id = ?
            """,
            (f"مرفوض من المدير: {payload.rejection_note}", now_iso(), task_id),
        )
        create_task_event(conn, task_id, "MANAGER_REJECTED", f"تم إرجاع الدليل للمالك المسؤول. الملاحظة: {payload.rejection_note}")
        conn.commit()
        updated = conn.execute("SELECT * FROM gap_tasks WHERE id = ?", (task_id,)).fetchone()
    return {"message": "تم إرجاع الدليل للمالك المسؤول", "task": row_to_dict(updated)}
